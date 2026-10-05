#!/usr/bin/env python3
"""Combina el listado oficial (JSON incrustado en las páginas de tiendas/restaurantes)
con las fichas extraídas (datos/multiplaza/locales.json) y genera:
- datos/multiplaza/directorio.json  (datos limpios por local, para la página web)
- documentos/Locales_Multiplaza_Escazu_Curridabat.xlsx (2 hojas, logos y datos)
"""
import json, re, html, os, io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW = f"{ROOT}/datos/multiplaza/raw"
MALLS = {"escazu": "Multiplaza Escazú", "curridabat": "Multiplaza Curridabat"}
INFO = {
 "escazu": {"direccion": "Autopista Próspero Fernández, San Rafael de Escazú, San José", "telefono": "+506 2201-6025",
            "horario": "L-S 10:00 a.m.–9:00 p.m. · D 10:00 a.m.–8:00 p.m.", "web": "https://multiplaza.com/escazu"},
 "curridabat": {"direccion": "Frente al Registro Nacional, Curridabat, San José", "telefono": "4001-7999 · multiplaza.curridabat@gruporoble.com",
            "horario": "L-S 10:00 a.m.–9:00 p.m. · D 10:00 a.m.–8:00 p.m.", "web": "https://multiplaza.com/curridabat"},
}
def norm(s):
    return re.sub(r"\s+", " ", (s or "")).strip()

def listado(mall, seccion):
    """Lee el arreglo allStores incrustado (Livewire) en la página de listado."""
    h = html.unescape(open(f"{RAW}/{mall}-{seccion}.html", encoding="utf-8").read())
    out = {}
    dec = json.JSONDecoder()
    for m in re.finditer(r'"allStores":\s*\[', h):
        try:
            arr, _ = dec.raw_decode(h, m.end() - 1)
        except Exception as e:
            print("allStores no parseable", mall, seccion, e); continue
        for d in arr:
            if d.get("slug"):
                out[d["slug"]] = {"slug": d["slug"], "name": d.get("store_name") or "", "logo": d.get("src_logo") or "",
                                  "addr": d.get("address") or "", "cat": d.get("category_name") or ""}
        if out:
            break
    return out

fichas = json.load(open(f"{ROOT}/datos/multiplaza/locales.json", encoding="utf-8"))
HOURS_RE = re.compile(r"(\b(?:L|Lun|Lunes|Mon|Dom|Domingo|Mar|Mi[eé]|Jue|Vie|S[aá]b)[^0-9]{0,12}\d{1,2}(?::\d{2})?\s?(?:AM|PM|am|pm|a\.m\.|p\.m\.)?.{0,60}?(?:AM|PM|am|pm|a\.m\.|p\.m\.|\d{1,2}:\d{2}))", re.I)

def limpiar_ficha(f, name, addr):
    t = norm(f.get("main_text", ""))
    # descripción: después del breadcrumb "Home Tiendas|Restaurantes <name>" y antes de la dirección
    desc = ""
    m = re.search(r"\bHome (?:Tiendas|Restaurantes|Stores|Restaurants|Kioscos) ", t)
    rest = t[m.end():] if m else t
    if name and rest.lower().startswith(name.lower()):
        rest = rest[len(name):].strip()
    cut = len(rest)
    a = norm(addr)
    ia = rest.find(a) if a else -1
    if ia >= 0:
        cut = ia
        after = rest[ia + len(a):]
    else:
        hm = HOURS_RE.search(rest)
        if hm:
            cut = hm.start(); after = rest[hm.start():]
        else:
            after = ""
    desc = rest[:cut].strip(" .") if cut > 0 else ""
    if desc and not desc.endswith((".", "!", "?")):
        desc += "."
    # horario: en `after`, antes del teléfono / E-mail / redes
    after = re.split(r"\s(?:\+?506[\s-]?)?\d{4}[\s-]?\d{4}\b|\sE-?mail\b|\sFacebook\b|\sInstagram\b|\sTik ?Tok\b|\sWhats[Aa]pp\b|\sSitio\b|\sWeb\b|\sWebsite\b", after, maxsplit=1)[0]
    after = re.split(r"\shttps?://", after, maxsplit=1)[0]
    horario = re.sub(r"^\s*Horario:?\s*", "", after, flags=re.I).strip(" ·-|,")
    if len(horario) > 90 or not re.search(r"\d", horario):
        hm = HOURS_RE.search(horario)
        horario = hm.group(1).strip() if hm else ""
    links = f.get("links", {})
    tel = [x for x in links.get("tel", []) if re.search(r"\d{4}", x)] or f.get("phones_in_text", []) or re.findall(r"\b[2-8]\d{7}\b", t)
    tel = sorted({re.sub(r"^(?:tel:)+", "", x).strip() for x in tel})
    tel = [re.sub(r"^(\d{4})(\d{4})$", r"\1-\2", x) for x in tel]
    webs = [w for w in links.get("web", []) if "cloudfront.net" not in w and "google.com/maps" not in w and "waze" not in w]
    return {
        "descripcion": desc[:900],
        "horario": horario,
        "telefono": " / ".join(tel),
        "correo": ", ".join(links.get("mailto", [])),
        "facebook": links.get("facebook", [""])[0] if links.get("facebook") else "",
        "instagram": links.get("instagram", [""])[0] if links.get("instagram") else "",
        "tiktok": links.get("tiktok", [""])[0] if links.get("tiktok") else "",
        "whatsapp": links.get("whatsapp", [""])[0] if links.get("whatsapp") else "",
        "web": webs[0] if webs else "",
        "logo_local": f.get("logo_local", ""),
    }

registros = []
for mall, mall_nombre in MALLS.items():
    fichas_mall = {f.get("slug"): f for f in fichas["malls"][mall]["locales"]}
    vistos = set()
    for seccion, tipo in (("tiendas", "T"), ("restaurantes", "R")):
        for slug, d in listado(mall, seccion).items():
            if slug in vistos:
                continue
            vistos.add(slug)
            f = fichas_mall.get(slug, {})
            extra = limpiar_ficha(f, d["name"], d["addr"]) if f else {}
            registros.append({
                "mall": mall, "mall_nombre": mall_nombre, "slug": slug, "nombre": norm(d["name"]),
                "categoria": norm(d["cat"]) or ("Restaurantes" if tipo == "R" else "Sin categoría"),
                "tipo": tipo, "ubicacion": norm(d["addr"]), "logo": "" if "sinlogo" in d["logo"] else d["logo"],
                "url": f"https://multiplaza.com/{mall}/{seccion}/{slug}",
                "descripcion": extra.get("descripcion", ""), "horario": extra.get("horario", ""),
                "telefono": extra.get("telefono", ""), "correo": extra.get("correo", ""),
                "facebook": extra.get("facebook", ""), "instagram": extra.get("instagram", ""),
                "tiktok": extra.get("tiktok", ""), "whatsapp": extra.get("whatsapp", ""), "web": extra.get("web", ""),
                "logo_local": extra.get("logo_local", ""),
                "gerente": "", "contacto_gerente": "", "notas": "",
            })
    # fichas que no estaban en el listado (por si acaso)
    for slug, f in fichas_mall.items():
        if slug not in vistos and "error" not in f:
            name = (f.get("title", "").split("|")[-1]).strip()
            extra = limpiar_ficha(f, name, "")
            registros.append({"mall": mall, "mall_nombre": mall_nombre, "slug": slug, "nombre": name,
                "categoria": "Restaurantes" if f.get("tipo_listado") == "restaurantes" else "Sin categoría",
                "tipo": "R" if f.get("tipo_listado") == "restaurantes" else "T", "ubicacion": "", "logo": f.get("logo_url", ""),
                "url": f.get("url", ""), **extra, "gerente": "", "contacto_gerente": "", "notas": ""})

registros.sort(key=lambda r: (r["mall"] != "escazu", r["tipo"], r["nombre"].lower()))
salida = {"generado": fichas["generado"], "fuente": "https://multiplaza.com (Grupo Roble)", "malls": {k: {"nombre": v, **INFO[k]} for k, v in MALLS.items()}, "locales": registros}
json.dump(salida, open(f"{ROOT}/datos/multiplaza/directorio.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)

# ---------- Excel ----------
F = "Arial"; VERDE = "0F3B2E"
thin = Side(style="thin", color="D9D2C3"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
pend = PatternFill("solid", fgColor="FFF6CC"); alt = PatternFill("solid", fgColor="F7F3EA")
cols = [("Logo", 12), ("Local", 28), ("Categoría", 26), ("Tipo", 16), ("Ubicación (etapa / nivel / local)", 38), ("Horario", 30),
        ("Teléfono", 15), ("Correo", 30), ("Facebook", 34), ("Instagram", 34), ("Otras redes / Web", 30),
        ("Gerente", 24), ("Contacto del gerente", 26), ("Notas", 24), ("Descripción", 60), ("Ficha oficial", 16)]
wb = Workbook(); wb.remove(wb.active)
tmpdir = f"{ROOT}/datos/multiplaza/_thumbs"; os.makedirs(tmpdir, exist_ok=True)
for mall, mall_nombre in MALLS.items():
    ws = wb.create_sheet(mall_nombre)
    rs = [r for r in registros if r["mall"] == mall]
    info = INFO[mall]
    ws["A1"] = f"Locales de {mall_nombre}"; ws["A1"].font = Font(name=F, bold=True, size=15, color=VERDE)
    ws["A2"] = f"{info['direccion']} · Tel. {info['telefono']} · {info['horario']} · {info['web']}"; ws["A2"].font = Font(name=F, size=9, color="6B6B66")
    ws["A3"] = f"Fuente: directorio oficial multiplaza.com (Grupo Roble), extraído el {fichas['generado']}. Las celdas en amarillo (Gerente, Contacto del gerente, Notas y datos vacíos) están por completar."; ws["A3"].font = Font(name=F, size=9, color="6B6B66")
    n = len(rs); fila0 = 7
    ws["A4"] = "Total de locales:"; ws["A4"].font = Font(name=F, bold=True); ws["B4"] = f"=COUNTA(B{fila0}:B{fila0+n-1})"; ws["B4"].font = Font(name=F, bold=True)
    ws["C4"] = "Tiendas y servicios:"; ws["C4"].font = Font(name=F, bold=True); ws["D4"] = f'=COUNTIF(D{fila0}:D{fila0+n-1},"Tienda / Servicio")'; ws["D4"].font = Font(name=F)
    ws["E4"] = "Restaurantes:"; ws["E4"].font = Font(name=F, bold=True); ws["F4"] = f'=COUNTIF(D{fila0}:D{fila0+n-1},"Restaurante")'; ws["F4"].font = Font(name=F)
    ws["G4"] = "Con teléfono:"; ws["G4"].font = Font(name=F, bold=True); ws["H4"] = f'=COUNTA(G{fila0}:G{fila0+n-1})'; ws["H4"].font = Font(name=F)
    ws["I4"] = "Gerentes por completar:"; ws["I4"].font = Font(name=F, bold=True); ws["J4"] = f'=COUNTBLANK(L{fila0}:L{fila0+n-1})'; ws["J4"].font = Font(name=F)
    for c, (h, w) in enumerate(cols, 1):
        cell = ws.cell(row=6, column=c, value=h); cell.font = Font(name=F, bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor=VERDE)
        cell.alignment = Alignment(vertical="center", wrap_text=True); cell.border = border
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[6].height = 30
    for i, r in enumerate(rs, start=fila0):
        tipo = "Restaurante" if r["tipo"] == "R" else "Tienda / Servicio"
        otras = " · ".join(x for x in (r["tiktok"], r["whatsapp"], r["web"]) if x)
        vals = ["", r["nombre"], r["categoria"], tipo, r["ubicacion"], r["horario"], r["telefono"], r["correo"], r["facebook"], r["instagram"], otras,
                r["gerente"], r["contacto_gerente"], r["notas"], r["descripcion"], "Ver ficha"]
        ws.row_dimensions[i].height = 42
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=c, value=v); cell.font = Font(name=F, size=10); cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if c in (6, 7, 8, 9, 10, 11, 12, 13, 14) and not v:
                cell.fill = pend
            elif i % 2 == 0:
                cell.fill = alt
            if c in (9, 10) and v:
                cell.hyperlink = v; cell.font = Font(name=F, size=10, color=VERDE, underline="single")
            if c == 16:
                cell.hyperlink = r["url"]; cell.font = Font(name=F, size=10, color=VERDE, underline="single")
        # logo
        lp = r.get("logo_local")
        if lp and "sinlogo" not in r.get("logo", "") and os.path.exists(f"{ROOT}/{lp}"):
            try:
                im = Image.open(f"{ROOT}/{lp}").convert("RGBA")
                bg = Image.new("RGBA", im.size, (255, 255, 255, 255)); bg.alpha_composite(im); im = bg.convert("RGB")
                im.thumbnail((76, 48))
                tp = f"{tmpdir}/{r['mall']}-{r['slug']}.png"; im.save(tp)
                xi = XLImage(tp); xi.anchor = f"A{i}"; ws.add_image(xi)
            except Exception as e:
                print("logo", r["slug"], e)
    ws.freeze_panes = "C7"; ws.auto_filter.ref = f"A6:P{fila0+n-1}"; ws.sheet_view.zoomScale = 90
wb.save(f"{ROOT}/documentos/Locales_Multiplaza_Escazu_Curridabat.xlsx")
print("locales:", {m: sum(1 for r in registros if r["mall"] == m) for m in MALLS})
print("con tel:", sum(1 for r in registros if r["telefono"]), "con correo:", sum(1 for r in registros if r["correo"]), "con fb:", sum(1 for r in registros if r["facebook"]), "con ig:", sum(1 for r in registros if r["instagram"]), "con horario:", sum(1 for r in registros if r["horario"]), "con desc:", sum(1 for r in registros if r["descripcion"]), "sin ficha:", sum(1 for r in registros if not r["logo_local"]))
